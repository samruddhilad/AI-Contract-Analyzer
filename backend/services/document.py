import csv
import io
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Dict, List

import pymupdf
from docx import Document

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".csv",
    ".xlsx",
    ".pptx",
    ".md",
    ".html",
    ".htm",
}


class HTMLTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.result: List[str] = []
        self.in_script = False

    def handle_starttag(self, tag: str, attrs: List[Any]):
        if tag.lower() in {"script", "style"}:
            self.in_script = True
        elif tag.lower() in {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr", "div", "br"}:
            self.result.append("\n")

    def handle_endtag(self, tag: str):
        if tag.lower() in {"script", "style"}:
            self.in_script = False
        elif tag.lower() in {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr"}:
            self.result.append("\n")

    def handle_data(self, data: str):
        if not self.in_script and data.strip():
            self.result.append(data.strip() + " ")

    def get_text(self) -> str:
        raw = "".join(self.result)
        raw = re.sub(r"[ \t]+", " ", raw)
        raw = re.sub(r"\n{3,}", "\n\n", raw)
        return raw.strip()


def extract_pdf_data(path: str) -> Dict[str, Any]:
    try:
        doc = pymupdf.open(path)
    except Exception as exc:
        raise ValueError(f"Corrupted or invalid PDF file: {exc}") from exc

    try:
        if doc.is_encrypted:
            if not doc.authenticate(""):
                raise ValueError("PDF document is encrypted or password-protected.")

        pages_data = []
        total_text_chars = 0
        has_images = False

        for idx, page in enumerate(doc):
            text = page.get_text("text").strip()
            total_text_chars += len(text)
            if not has_images and len(page.get_images()) > 0:
                has_images = True
            pages_data.append({
                "page": idx + 1,
                "source": f"Page {idx + 1}",
                "text": text,
            })

        is_scanned = total_text_chars < 50 and has_images
        limitation_note = ""
        if is_scanned:
            limitation_note = (
                "[Notice: Scanned document detected with minimal digital text layer. "
                "Extraction limitation disclosed; OCR text may be incomplete.]"
            )

        full_text_parts = []
        if limitation_note:
            full_text_parts.append(limitation_note)

        for p in pages_data:
            if p["text"]:
                full_text_parts.append(f"--- [Page {p['page']}] ---\n{p['text']}")

        full_text = "\n\n".join(full_text_parts).strip()
        return {
            "text": full_text,
            "pages": pages_data,
            "metadata": {
                "format": "pdf",
                "page_count": len(doc),
                "is_scanned": is_scanned,
                "limitation_note": limitation_note,
            },
        }
    finally:
        doc.close()


def extract_docx_data(path: str) -> Dict[str, Any]:
    try:
        doc = Document(path)
    except Exception as exc:
        raise ValueError(f"Corrupted or invalid DOCX document: {exc}") from exc

    sections = []
    current_section_title = "Document Body"
    current_section_lines = []

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if p.style and p.style.name and p.style.name.startswith("Heading"):
            if current_section_lines:
                sections.append({
                    "source": current_section_title,
                    "text": "\n".join(current_section_lines),
                })
                current_section_lines = []
            current_section_title = txt
        else:
            current_section_lines.append(txt)

    if current_section_lines:
        sections.append({
            "source": current_section_title,
            "text": "\n".join(current_section_lines),
        })

    table_rows = []
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                table_rows.append(row_text)

    if table_rows:
        sections.append({
            "source": "Document Tables",
            "text": "\n".join(table_rows),
        })

    full_text = "\n\n".join(f"[{s['source']}]\n{s['text']}" for s in sections if s["text"]).strip()
    return {
        "text": full_text,
        "pages": [{"page": i + 1, "source": s["source"], "text": s["text"]} for i, s in enumerate(sections)],
        "metadata": {
            "format": "docx",
            "section_count": len(sections),
            "is_scanned": False,
            "limitation_note": "",
        },
    }


def extract_txt_data(path: str) -> Dict[str, Any]:
    raw_bytes = Path(path).read_bytes()
    if not raw_bytes:
        raise ValueError("Uploaded file is empty (0 bytes).")

    text = None
    for enc in ["utf-8", "utf-8-sig", "utf-16", "latin-1", "cp1252"]:
        try:
            text = raw_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        text = raw_bytes.decode("utf-8", errors="replace")

    text = text.strip()
    return {
        "text": text,
        "pages": [{"page": 1, "source": "Text Document", "text": text}],
        "metadata": {"format": "txt", "is_scanned": False, "limitation_note": ""},
    }


def extract_csv_data(path: str) -> Dict[str, Any]:
    raw_bytes = Path(path).read_bytes()
    if not raw_bytes:
        raise ValueError("Uploaded CSV file is empty (0 bytes).")

    decoded = None
    for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            decoded = raw_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not decoded:
        decoded = raw_bytes.decode("utf-8", errors="replace")

    try:
        reader = csv.reader(io.StringIO(decoded))
        rows = list(reader)
    except Exception as exc:
        raise ValueError(f"Corrupted or invalid CSV file: {exc}") from exc

    if not rows:
        raise ValueError("CSV file contains no readable rows.")

    formatted_lines = []
    headers = rows[0] if rows else []
    formatted_lines.append("CSV Columns: " + " | ".join(headers))
    for r in rows[1:500]:
        if any(cell.strip() for cell in r):
            formatted_lines.append(" | ".join(cell.strip() for cell in r))

    text = "\n".join(formatted_lines).strip()
    return {
        "text": text,
        "pages": [{"page": 1, "source": "CSV Data Table", "text": text}],
        "metadata": {
            "format": "csv",
            "row_count": len(rows),
            "col_count": len(headers),
            "is_scanned": False,
            "limitation_note": f"Analyzed {min(len(rows), 500)} of {len(rows)} rows.",
        },
    }


def extract_xlsx_data(path: str) -> Dict[str, Any]:
    if openpyxl is None:
        raise ValueError("openpyxl is not installed. Please install openpyxl to support Excel spreadsheets.")

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"Corrupted or invalid XLSX spreadsheet: {exc}") from exc

    pages = []
    full_text_parts = []

    try:
        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            sheet_rows = []
            for row in ws.iter_rows(values_only=True):
                non_empty = [str(val).strip() for val in row if val is not None and str(val).strip()]
                if non_empty:
                    sheet_rows.append(" | ".join(non_empty))
                if len(sheet_rows) >= 300:
                    break

            if sheet_rows:
                sheet_content = f"--- [Sheet: {sheet_name}] ---\n" + "\n".join(sheet_rows)
                pages.append({"page": sheet_idx + 1, "source": f"Sheet: {sheet_name}", "text": sheet_content})
                full_text_parts.append(sheet_content)

        full_text = "\n\n".join(full_text_parts).strip()
        if not full_text:
            raise ValueError("Spreadsheet does not contain readable cell data.")

        return {
            "text": full_text,
            "pages": pages,
            "metadata": {
                "format": "xlsx",
                "sheet_count": len(wb.sheetnames),
                "is_scanned": False,
                "limitation_note": "",
            },
        }
    finally:
        wb.close()


def extract_pptx_data(path: str) -> Dict[str, Any]:
    if Presentation is None:
        raise ValueError("python-pptx is not installed. Please install python-pptx to support PowerPoint presentations.")

    try:
        prs = Presentation(path)
    except Exception as exc:
        raise ValueError(f"Corrupted or invalid PPTX presentation: {exc}") from exc

    pages = []
    full_text_parts = []

    for slide_idx, slide in enumerate(prs.slides):
        slide_texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_texts.append(shape.text.strip())
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_txt:
                        slide_texts.append(row_txt)

        slide_content = f"--- [Slide {slide_idx + 1}] ---\n" + "\n".join(slide_texts)
        pages.append({"page": slide_idx + 1, "source": f"Slide {slide_idx + 1}", "text": slide_content})
        if slide_texts:
            full_text_parts.append(slide_content)

    full_text = "\n\n".join(full_text_parts).strip()
    return {
        "text": full_text,
        "pages": pages,
        "metadata": {
            "format": "pptx",
            "slide_count": len(prs.slides),
            "is_scanned": False,
            "limitation_note": "",
        },
    }


def extract_html_data(path: str) -> Dict[str, Any]:
    raw_bytes = Path(path).read_bytes()
    if not raw_bytes:
        raise ValueError("Uploaded HTML file is empty (0 bytes).")

    decoded = None
    for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            decoded = raw_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not decoded:
        decoded = raw_bytes.decode("utf-8", errors="replace")

    parser = HTMLTextExtractor()
    parser.feed(decoded)
    text = parser.get_text()

    if not text:
        raise ValueError("HTML document contains no readable text content.")

    return {
        "text": text,
        "pages": [{"page": 1, "source": "HTML Document", "text": text}],
        "metadata": {"format": "html", "is_scanned": False, "limitation_note": ""},
    }


def extract_document(path: str) -> Dict[str, Any]:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Document file not found at: {path}")

    if p.stat().st_size == 0:
        raise ValueError("Uploaded file is empty (0 bytes).")

    suffix = p.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: '{suffix}'. Supported formats: PDF, DOCX, TXT, CSV, XLSX, PPTX, MD, HTML."
        )

    if suffix == ".pdf":
        data = extract_pdf_data(path)
    elif suffix == ".docx":
        data = extract_docx_data(path)
    elif suffix in {".txt", ".md"}:
        data = extract_txt_data(path)
        data["metadata"]["format"] = suffix.replace(".", "")
    elif suffix == ".csv":
        data = extract_csv_data(path)
    elif suffix == ".xlsx":
        data = extract_xlsx_data(path)
    elif suffix == ".pptx":
        data = extract_pptx_data(path)
    elif suffix in {".html", ".htm"}:
        data = extract_html_data(path)
    else:
        raise ValueError(f"Unsupported file type: '{suffix}'.")

    if not data["text"].strip():
        raise ValueError("No readable text could be extracted from this document.")

    return data


def extract_text(path: str) -> str:
    """Backward-compatible helper returning cleaned extracted string."""
    data = extract_document(path)
    return data["text"]
